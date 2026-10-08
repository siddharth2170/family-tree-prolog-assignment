from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Family_Tree_Program_APA7.docx"
CODE = (ROOT / "family_tree.pl").read_text()


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def add_page_number(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_font(run, size=12)


def add_centered(text="", bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, bold=bold)
    return p


def add_body(text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    set_font(r)
    return p


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_font(r, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.5)
add_page_number(section)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)

# Some Word/LibreOffice installations apply a theme border to the built-in
# Title style. Remove it so the APA title is separated by whitespace only.
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

# APA student title page
for _ in range(3):
    add_centered()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.line_spacing = 2
title.paragraph_format.space_after = Pt(0)
r = title.add_run("Family Tree Program in Prolog")
set_font(r, size=12, bold=True)
add_centered("Student Name")
add_centered("Department and University")
add_centered("Course Number and Course Name")
add_centered("Instructor Name")
add_centered("October 8, 2026")

doc.add_page_break()
add_heading("Family Tree Program in Prolog")
add_body(
    "This project implements a family knowledge base in Prolog. The program stores direct relationships as parent/2, male/1, and female/1 facts and derives other relationships through rules. Separating facts from rules makes the tree easy to extend without changing its reasoning structure. The sample data spans three generations so the required child, sibling, grandparent, cousin, and descendant queries all produce meaningful results."
)
add_heading("Rule Design", level=2)
add_body(
    "The child rule reverses parent/2 so a query can begin with a parent and return each child. The grandparent rule joins two parent facts through an intermediate person. The sibling rule requires distinct people to share a parent, and once/1 prevents duplicate success when full siblings share two parents. The cousin rule composes these relationships by testing whether the two people's parents are siblings. Each predicate therefore expresses one relationship and can be reused by later rules."
)
add_heading("Recursion and Logical Inference", level=2)
add_body(
    "The descendant rule has a base case for a direct child and a recursive case that follows one parent link before testing the remaining path. Prolog's backtracking explores each branch to find children, grandchildren, and later generations. The base case returns nearby descendants first, and each recursive step moves downward through the finite tree."
)
add_heading("Challenges and Verification", level=2)
add_body(
    "The main challenge was duplicate answers: two shared parents can satisfy sibling/2 twice, and composed rules may expose multiple proof paths. The program uses once/1 to limit sibling proofs, setof/3 to return unique sorted query results, and dif/2 to prevent self-sibling and self-cousin answers."
)

repo_heading = add_heading("Repository Link")
repo_heading.paragraph_format.page_break_before = True
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 2
p.paragraph_format.space_after = Pt(0)
r = p.add_run("GitHub repository: ")
set_font(r, bold=True)
r = p.add_run("https://github.com/siddharth2170/family-tree-prolog-assignment")
set_font(r)

add_heading("Appendix A")
add_centered("Prolog Source Code", bold=True)
for line in CODE.rstrip().splitlines():
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    r = p.add_run(line if line else " ")
    set_font(r, name="Courier New", size=8.5)

doc.add_page_break()
add_heading("Appendix B")
add_centered("Sample Queries and Expected Output", bold=True)
intro = doc.add_paragraph()
intro.paragraph_format.line_spacing = 2
intro.paragraph_format.space_after = Pt(6)
r = intro.add_run("Load the file with ")
set_font(r)
r = intro.add_run("swipl -s family_tree.pl")
set_font(r, name="Courier New", size=10)
r = intro.add_run(". At the Prolog prompt, enter each query exactly as shown.")
set_font(r)

queries = [
    ("Children", "?- setof(Child, child(Child, john), Children).", "Children = [mary, robert]."),
    ("Siblings", "?- setof(Sibling, sibling(mary, Sibling), Siblings).", "Siblings = [robert]."),
    ("Grandchildren", "?- setof(Person, grandparent(george, Person), Grandchildren).", "Grandchildren = [emily, mary, michael, robert]."),
    ("Cousin test", "?- cousin(mary, emily).", "true."),
    ("Negative cousin test", "?- cousin(mary, robert).", "false."),
    ("Cousins", "?- setof(Cousin, cousin(mary, Cousin), Cousins).", "Cousins = [emily, michael]."),
    ("Recursive descendants", "?- setof(Person, descendant(Person, george), Descendants).", "Descendants = [emily, john, liam, mary, michael, robert, sophia, susan]."),
    ("Ancestor test", "?- descendant(sophia, george).", "true."),
]

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.2)
table.columns[1].width = Inches(3.35)
table.columns[2].width = Inches(1.95)
set_table_borders(table)
headers = ("Purpose", "Query", "Expected output")
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = table.columns[i].width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, name="Arial", size=9, bold=True)
    r.font.color.rgb = RGBColor(255, 255, 255)

for idx, (purpose, query, output) in enumerate(queries):
    cells = table.add_row().cells
    if idx % 2:
        for cell in cells:
            set_cell_shading(cell, "EAF2F8")
    values = (purpose, query, output)
    for i, value in enumerate(values):
        cells[i].width = table.columns[i].width
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
        p = cells[i].paragraphs[0]
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(value)
        set_font(r, name="Arial" if i == 0 else "Courier New", size=8.5)

note = doc.add_paragraph()
note.paragraph_format.line_spacing = 2
note.paragraph_format.space_before = Pt(8)
note.paragraph_format.space_after = Pt(0)
r = note.add_run("Note. ")
set_font(r, italic=True)
r = note.add_run("setof/3 removes duplicate solutions and displays the results as a sorted list. In an interactive query without setof/3, enter a semicolon to request the next solution.")
set_font(r)

doc.core_properties.title = "Family Tree Program in Prolog"
doc.core_properties.subject = "Prolog family relationships and recursive inference"
doc.core_properties.author = "Student Name"
doc.core_properties.keywords = "Prolog, family tree, recursion, logical inference, APA 7"
doc.save(OUT)
print(OUT)
